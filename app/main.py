import asyncio
import base64
import json
import math
import os
import re
import tempfile
import traceback
import uuid
from copy import deepcopy
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import httpx
import pandas as pd
from fastapi import FastAPI, HTTPException
from fastapi.responses import FileResponse, JSONResponse
from pydantic import BaseModel, Field
from pptx import Presentation

APP_VERSION = "13.1.1-ux-polish"

MISSING_PPT_VALUE = "N/A"
MISSING_DISPLAY_VALUE = "N/A (not specified in source file)"
EXEC_SUMMARY_MAX_CHARS = 550

PPTX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
MAX_RETURN_FILE_BYTES = 10 * 1024 * 1024

EMU_PER_PX = 9525
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

app = FastAPI(title="PCA Automation Generator", version=APP_VERSION)

BASE_DIR = Path(__file__).resolve().parent
ROOT_DIR = BASE_DIR.parent


def resolve_existing_path(env_key: str, candidates: List[Path]) -> Path:
    env_value = os.getenv(env_key)
    if env_value:
        return Path(env_value)

    for candidate in candidates:
        if candidate.exists():
            return candidate

    # Return the first candidate so health output still makes the expected path clear.
    return candidates[0]


MODULAR_TEMPLATE_PATH = resolve_existing_path(
    "MODULAR_TEMPLATE_PATH",
    [
        BASE_DIR / "templates" / "modular_sections_master.pptx",
        ROOT_DIR / "templates" / "modular_sections_master.pptx",
    ],
)

SLIDE_DECK_TEMPLATE_PATH = resolve_existing_path(
    "SLIDE_DECK_TEMPLATE_PATH",
    [
        BASE_DIR / "templates" / "slide_deck_master.pptx",
        ROOT_DIR / "templates" / "slide_deck_master.pptx",
    ],
)

def template_path_for(base_path: Path, language: str) -> Path:
    """Return the language-specific master (e.g. *_ja.pptx) if it exists,
    otherwise fall back to the base English master."""
    if language and language.lower() == "ja":
        ja = base_path.with_name(f"{base_path.stem}_ja{base_path.suffix}")
        if ja.exists():
            return ja
    return base_path


SECTION_REGISTRY_PATH = resolve_existing_path(
    "SECTION_REGISTRY_PATH",
    [
        BASE_DIR / "section_registry" / "section_registry.json",
        ROOT_DIR / "app" / "section_registry" / "section_registry.json",
        ROOT_DIR / "section_registry" / "section_registry.json",
    ],
)

RULES_MASTER_PATH = Path(
    os.getenv("RULES_MASTER_PATH", ROOT_DIR / "assets" / "PCA_GPT_Rules_Master.xlsx")
)

PUBLIC_BASE_URL = os.getenv(
    "PUBLIC_BASE_URL",
    "https://pca-modular-builder-v12.onrender.com"
)

GENERATED_FILES_DIR = Path(
    os.getenv("GENERATED_FILES_DIR", tempfile.gettempdir())
) / "pca_modular_generated"

GENERATED_FILES_DIR.mkdir(parents=True, exist_ok=True)


class OpenAIFileRef(BaseModel):
    id: Optional[str] = None
    name: Optional[str] = None
    mime_type: Optional[str] = None
    download_link: Optional[str] = None


class FileRefsPayload(BaseModel):
    openaiFileIdRefs: List[OpenAIFileRef] = Field(...)
    exec_summary: Optional[str] = None


class ModularOnePagerRequest(BaseModel):
    selected_sections: List[str] = Field(default_factory=list)
    placeholder_values: Optional[Dict[str, Any]] = Field(default_factory=dict)
    top_margin_px: int = 100
    bottom_margin_px: int = 100
    section_spacing_px: int = 60


class ModularOnePagerFromEocRequest(BaseModel):
    openaiFileIdRefs: List[OpenAIFileRef] = Field(...)
    selected_sections: List[str] = Field(default_factory=list)
    exec_summary: Optional[str] = None
    top_margin_px: int = 100
    bottom_margin_px: int = 100
    section_spacing_px: int = 60


class SlideDeckFromEocRequest(BaseModel):
    openaiFileIdRefs: List[OpenAIFileRef] = Field(...)
    selected_sections: List[str] = Field(default_factory=list)
    one_pager_sections: List[str] = Field(default_factory=list)
    custom_deck_sections: List[str] = Field(default_factory=list)
    deck_mode: str = "matching"
    exec_summary: Optional[str] = None


@app.get("/")
def read_root():
    return {
        "status": "healthy",
        "message": "PCA Automation Generator API is running successfully.",
        "mcp_path": "/mcp"
    }


@app.get("/health")
def health():
    return {
        "status": "ok",
        "app_version": APP_VERSION,
        "modular_template_exists": MODULAR_TEMPLATE_PATH.exists(),
        "slide_deck_template_exists": SLIDE_DECK_TEMPLATE_PATH.exists(),
        "section_registry_exists": SECTION_REGISTRY_PATH.exists(),
        "rules_master_exists": RULES_MASTER_PATH.exists(),
        "generated_files_dir": str(GENERATED_FILES_DIR),
    }


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except Exception:
        pass
    return re.sub(r"\s+", " ", str(value)).strip()


def normalize_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", " ", clean_text(value).lower()).strip()


def normalize_key(value: Any) -> str:
    return normalize_header(value).replace(" ", "_")


def normalize_text(value: Any) -> str:
    return re.sub(r"[^a-z0-9%./()\- ]+", " ", clean_text(value).lower()).strip()


def clean_dimension_label(value: Any) -> str:
    value = clean_text(value)
    value = re.sub(r"\bMISCRL\b", "", value, flags=re.IGNORECASE)
    value = re.sub(r"([a-z])([A-Z])", r"\1 \2", value)
    value = re.sub(r"\bDisplay\s*Animated\b", "Display Animated", value, flags=re.IGNORECASE)
    value = re.sub(r"\bIn\s*Situ\s*Video\b", "In Situ Video", value, flags=re.IGNORECASE)
    value = re.sub(r"[-_]+", " - ", value)
    value = re.sub(r"\s+-\s+-\s+", " - ", value)
    value = re.sub(r"\s{2,}", " ", value)
    return value.strip(" -")


def safe_number(value: Any) -> Optional[float]:
    try:
        if value is None:
            return None
        if isinstance(value, str):
            value = (
                value.replace(",", "")
                .replace("£", "")
                .replace("$", "")
                .replace("€", "")
                .replace("%", "")
                .strip()
            )
        num = float(value)
        if math.isnan(num) or math.isinf(num):
            return None
        return num
    except Exception:
        return None


def abs_number(value: Any) -> Optional[float]:
    num = safe_number(value)
    return abs(num) if num is not None else None


def format_number(value: Optional[float], decimals: int = 0) -> Optional[str]:
    if value is None:
        return None
    if decimals == 0:
        return f"{int(round(value)):,}"
    return f"{value:,.{decimals}f}"


def format_currency(value: Optional[float], symbol: str = "€") -> Optional[str]:
    if value is None:
        return None
    return f"{symbol}{value:,.2f}"


def format_percent(value: Any) -> Optional[str]:
    if isinstance(value, str) and "%" in value:
        return clean_text(value)

    num = safe_number(value)
    if num is None:
        return None

    if num <= 1:
        num = num * 100

    return f"{round(num, 2)}%"


def parse_date_from_any(value: Any) -> Optional[datetime]:
    if value is None:
        return None

    if isinstance(value, datetime):
        return value

    try:
        if pd.isna(value):
            return None
    except Exception:
        pass

    text = clean_text(value)
    if not text:
        return None

    try:
        dt = pd.to_datetime(text, errors="coerce")
        if pd.notna(dt):
            py_dt = dt.to_pydatetime()
            if 2000 <= py_dt.year <= 2100:
                return py_dt
    except Exception:
        pass

    match = re.search(r"\b(20\d{2}-\d{2}-\d{2})\b", text)
    if match:
        try:
            return datetime.strptime(match.group(1), "%Y-%m-%d")
        except Exception:
            pass

    match = re.search(r"\b(\d{1,2}[/-]\d{1,2}[/-]20\d{2})\b", text)
    if match:
        try:
            dt = pd.to_datetime(match.group(1), dayfirst=True, errors="coerce")
            if pd.notna(dt):
                return dt.to_pydatetime()
        except Exception:
            pass

    return None


def format_date_short(start_dt: Optional[datetime], end_dt: Optional[datetime]) -> Optional[str]:
    if start_dt and end_dt:
        return f"{start_dt.strftime('%d %b')} - {end_dt.strftime('%d %b')}"
    return None


def format_date_full(start_dt: Optional[datetime], end_dt: Optional[datetime]) -> Optional[str]:
    if start_dt and end_dt:
        if start_dt.year == end_dt.year:
            return f"{start_dt.strftime('%d %B')} - {end_dt.strftime('%d %B %Y')}"
        return f"{start_dt.strftime('%d %B %Y')} - {end_dt.strftime('%d %B %Y')}"
    return None


def format_quarter_from_date(end_dt: Optional[datetime]) -> Optional[str]:
    if not end_dt:
        return None
    q = ((end_dt.month - 1) // 3) + 1
    return f"Q{q} {end_dt.year}"


def fill_missing_for_ppt(mapped: Dict[str, Any]) -> Dict[str, Any]:
    for key, value in mapped.items():
        if value is None or clean_text(value) == "":
            mapped[key] = MISSING_PPT_VALUE
    return mapped


def build_display_values(mapped: Dict[str, Any]) -> Dict[str, Any]:
    return {
        key: MISSING_DISPLAY_VALUE if value == MISSING_PPT_VALUE else value
        for key, value in mapped.items()
    }


def first_non_empty(series: pd.Series) -> Any:
    for value in series:
        if clean_text(value) != "":
            return value
    return None


def is_blank_row(values: List[Any]) -> bool:
    return all(clean_text(v) == "" for v in values)


MARKET_CODE_MAP = {
    "france": "FR", "french": "FR", "fr": "FR",
    "united kingdom": "UK", "uk": "UK", "gb": "UK", "great britain": "UK", "britain": "UK",
    "italy": "IT", "italian": "IT", "it": "IT",
    "spain": "ES", "spanish": "ES", "es": "ES",
    "germany": "DE", "german": "DE", "de": "DE",
    "netherlands": "NL", "dutch": "NL", "nl": "NL",
    "belgium": "BE", "be": "BE",
    "switzerland": "CH", "ch": "CH",
    "austria": "AT", "at": "AT",
    "portugal": "PT", "pt": "PT",
    "ireland": "IE", "ie": "IE",
    "usa": "US", "us": "US", "united states": "US", "united states of america": "US",
    "canada": "CA", "ca": "CA",
    "australia": "AU", "au": "AU",
    "japan": "JP", "jp": "JP",
    "hong kong": "HK", "hk": "HK",
    "singapore": "SG", "sg": "SG",
    "thailand": "TH", "thai": "TH", "th": "TH",
    "taiwan": "TW", "tw": "TW",
    "china": "CN", "cn": "CN",
    "korea": "KR", "south korea": "KR", "kr": "KR",
    "india": "IN", "in": "IN",
    "mexico": "MX", "mx": "MX",
    "brazil": "BR", "br": "BR",
}


def normalise_market_label(value: Any) -> str:
    raw = clean_dimension_label(value)
    if not raw:
        return ""

    raw = re.sub(
        r"\b(total|totals|overall|summary|market|markets|country|countries)\b",
        "",
        raw,
        flags=re.IGNORECASE,
    )
    raw = raw.strip(" ,-")

    if not raw:
        return ""

    parts = re.split(r"[,;/|]+", raw)
    out = []

    for part in parts:
        part = clean_text(part).strip(" -")
        if not part:
            continue

        key = part.lower()
        code = MARKET_CODE_MAP.get(key)

        if not code:
            if re.fullmatch(r"[A-Za-z]{2,3}", part):
                code = part.upper()
            else:
                code = part.upper() if len(part) <= 3 else part

        if code and code not in out:
            out.append(code)

    return ", ".join(out)


ALIASES = {
    "campaign": ["campaign", "campaign name"],
    "client_brand": ["client", "brand", "advertiser", "brand / client"],
    "impressions": ["impressions", "delivered impressions", "served impressions"],
    "ctr": ["ctr", "click through rate", "click-through rate"],
    "engagement_rate": ["engagement rate", "engagement %", "er", "total er", "overall er"],
    "vcr": ["video completion rate", "vcr", "completed view rate", "video % complete"],
    "on_screen": ["mobkoi on screen", "on screen", "on-screen", "mrc viewability", "viewability", "viewable rate"],
    "mobkoi_on_screen": ["mobkoi on screen"],
    "mrc_viewability": ["mrc viewability"],
    "spend": ["actual spend", "spend", "total spend", "media spend"],
    "sold_paid_units": ["sold paid units"],
    "delivered_overall_av_units": ["delivered overall av units", "delivered overall av (units)"],
    "delivery_incl_av": ["delivery percentage incl av", "delivery percentage (incl av)"],
    "delivered_av_amount": ["delivered av amount currency", "delivered av amount (currency)", "worth of added value"],
    "site": ["site", "publisher", "domain", "environment", "property", "placement", "title", "inventory", "app", "website"],
    "geo": ["geo", "market", "markets", "country", "countries", "region", "territory", "location", "locale"],
    "format": ["format", "formats", "creative", "creative format", "ad format", "unit type", "product", "product type", "placement type"],
    "date": ["date", "report date", "served date", "live date"],
}


def header_match_score(header: str, aliases: List[str]) -> int:
    norm = normalize_header(header)
    best = 0

    for alias in aliases:
        alias_norm = normalize_header(alias)
        if norm == alias_norm:
            best = max(best, 100)
        elif alias_norm in norm:
            best = max(best, 85)
        elif norm in alias_norm:
            best = max(best, 60)

    return best


def find_col(df: pd.DataFrame, key: str) -> Optional[str]:
    aliases = ALIASES.get(key, [key])
    best_col = None
    best_score = 0

    for col in df.columns:
        score = header_match_score(str(col), aliases)
        if score > best_score:
            best_score = score
            best_col = col

    return best_col if best_score >= 60 else None


def score_header_row(row_values: List[Any]) -> int:
    keywords = [
        "campaign", "impressions", "ctr", "engagement", "vcr", "completion",
        "spend", "site", "publisher", "domain", "environment", "property",
        "placement", "title", "inventory", "app", "website",
        "geo", "market", "markets", "country", "countries", "region", "territory", "location", "locale",
        "format", "formats", "creative", "product", "unit type",
        "date", "report date", "served date", "live date",
        "sold paid units", "delivered overall av", "delivery percentage", "mrc viewability",
    ]

    score = 0
    for cell in row_values:
        text = normalize_text(cell)
        for kw in keywords:
            if kw in text:
                score += 1

    return score


def find_candidate_header_rows(df: pd.DataFrame) -> List[int]:
    candidates = []

    for i in range(min(120, len(df))):
        score = score_header_row(df.iloc[i].tolist())
        if score >= 3:
            candidates.append(i)

    deduped = []
    for idx in candidates:
        if not deduped or idx - deduped[-1] > 2:
            deduped.append(idx)

    return deduped


def build_block_from_header(df: pd.DataFrame, header_row: int) -> pd.DataFrame:
    if df is None or df.empty:
        return pd.DataFrame()

    temp = df.iloc[header_row:].copy()
    if temp.empty:
        return pd.DataFrame()

    temp.columns = [clean_text(c) for c in temp.iloc[0]]
    temp = temp[1:].reset_index(drop=True)
    temp = temp.dropna(axis=1, how="all")

    rows = []

    for _, row in temp.iterrows():
        row_vals = row.tolist()
        if is_blank_row(row_vals):
            break
        rows.append(row_vals)

    if not rows:
        return pd.DataFrame(columns=temp.columns)

    block = pd.DataFrame(rows, columns=temp.columns)

    header_norm = [normalize_header(c) for c in block.columns]
    keep_rows = []

    for _, row in block.iterrows():
        row_norm = [normalize_header(v) for v in row.tolist()]
        same_count = 0

        for a, b in zip(header_norm[:12], row_norm[:12]):
            if a and b and a == b:
                same_count += 1

        keep_rows.append(same_count < 3)

    return block[pd.Series(keep_rows).values].reset_index(drop=True)


def classify_table(df: pd.DataFrame) -> str:
    if df is None or df.empty:
        return "unknown"

    site_col = find_col(df, "site")
    geo_col = find_col(df, "geo")
    format_col = find_col(df, "format")

    has_ctr = find_col(df, "ctr") is not None
    has_er = find_col(df, "engagement_rate") is not None
    has_vcr = find_col(df, "vcr") is not None
    has_impressions = find_col(df, "impressions") is not None

    delivery_hits = sum([
        1 if find_col(df, "sold_paid_units") else 0,
        1 if find_col(df, "delivered_overall_av_units") else 0,
        1 if find_col(df, "delivery_incl_av") else 0,
        1 if find_col(df, "delivered_av_amount") else 0,
    ])

    # Site performance tables often also contain geo/global_local columns.
    # Prioritise Site classification when a site/publisher column appears with performance metrics.
    if site_col and (has_ctr or has_er or has_vcr):
        return "site"

    if delivery_hits >= 2:
        return "campaign_delivery"

    if geo_col and has_impressions and not site_col:
        return "geo"

    if format_col and has_impressions and not site_col:
        return "format"

    date_col = find_col(df, "date")
    if date_col:
        col_data = df[date_col]
        if isinstance(col_data, pd.DataFrame):
            col_data = col_data.iloc[:, 0]
        sample = [parse_date_from_any(v) for v in col_data.head(10).values]
        if any(v is not None for v in sample):
            return "date"

    kpi_hits = sum([
        1 if find_col(df, "campaign") else 0,
        1 if has_impressions else 0,
        1 if has_ctr else 0,
        1 if has_er else 0,
        1 if has_vcr else 0,
        1 if (find_col(df, "mobkoi_on_screen") or find_col(df, "mrc_viewability") or find_col(df, "on_screen")) else 0,
        1 if find_col(df, "spend") else 0,
    ])

    if kpi_hits >= 3:
        return "campaign_kpi_summary"

    if site_col:
        return "site"
    if geo_col:
        return "geo"
    if format_col:
        return "format"

    return "unknown"


def table_quality_score(df: pd.DataFrame, table_type: str) -> int:
    if df is None or df.empty:
        return -999

    score = min(len(df.columns), 20) + min(len(df), 10)

    weights = {
        "campaign_kpi_summary": ["campaign", "impressions", "ctr", "engagement_rate", "vcr", "spend"],
        "campaign_delivery": ["sold_paid_units", "delivered_overall_av_units", "delivery_incl_av", "delivered_av_amount"],
        "site": ["site", "ctr", "engagement_rate", "vcr"],
        "geo": ["geo", "impressions"],
        "format": ["format", "impressions"],
        "date": ["date"],
    }

    for key in weights.get(table_type, []):
        if find_col(df, key):
            score += 12 if table_type == "campaign_delivery" else 10

    # Strongly favour true Site tables that include publisher/title plus multiple performance metrics.
    # This prevents Site+Geo tables being selected as generic geo tables and losing top-publisher output.
    if table_type == "site":
        metric_hits = sum([
            1 if find_col(df, "ctr") else 0,
            1 if find_col(df, "engagement_rate") else 0,
            1 if find_col(df, "vcr") else 0,
            1 if find_col(df, "impressions") else 0,
        ])
        if find_col(df, "site"):
            score += 30
        score += metric_hits * 15

    if table_type == "campaign_kpi_summary":
        if find_col(df, "mobkoi_on_screen"):
            score += 12
        elif find_col(df, "mrc_viewability"):
            score += 6

    return score


def detect_tables(sheets: Dict[str, pd.DataFrame]) -> Dict[str, pd.DataFrame]:
    candidates_by_type: Dict[str, List[Tuple[int, pd.DataFrame]]] = {
        "campaign_kpi_summary": [],
        "campaign_delivery": [],
        "site": [],
        "geo": [],
        "format": [],
        "date": [],
    }

    for _, raw_df in sheets.items():
        if raw_df is None or raw_df.empty:
            continue

        for header_row in find_candidate_header_rows(raw_df):
            block = build_block_from_header(raw_df, header_row)
            if block.empty:
                continue

            table_type = classify_table(block)
            if table_type == "unknown":
                continue

            score = table_quality_score(block, table_type)
            candidates_by_type[table_type].append((score, block))

    detected: Dict[str, pd.DataFrame] = {}

    for table_type, candidates in candidates_by_type.items():
        if candidates:
            candidates.sort(key=lambda x: x[0], reverse=True)
            detected[table_type] = candidates[0][1]

    return detected


def scan_workbook_texts(sheets: Dict[str, pd.DataFrame]) -> List[str]:
    texts = []

    for df in sheets.values():
        if df is None or df.empty:
            continue

        for row in df.head(20).values:
            for cell in row:
                text = clean_text(cell)
                if text:
                    texts.append(text)

    return texts


def extract_dates(
    sheets: Dict[str, pd.DataFrame],
    detected: Dict[str, pd.DataFrame],
    filename: str = "",
) -> Tuple[Optional[datetime], Optional[datetime]]:
    date_df = detected.get("date")

    if date_df is not None:
        date_col = find_col(date_df, "date")
        if date_col:
            col_data = date_df[date_col]
            if isinstance(col_data, pd.DataFrame):
                col_data = col_data.iloc[:, 0]

            vals = [parse_date_from_any(v) for v in col_data.tolist()]
            vals = [v for v in vals if v]
            if vals:
                return min(vals), max(vals)

    for text in scan_workbook_texts(sheets):
        if "report date" in normalize_text(text):
            dt = parse_date_from_any(text)
            if dt:
                return dt - timedelta(days=6), dt

    match = re.search(r"\b(20\d{2}-\d{2}-\d{2})\b", filename or "")
    if match:
        end_dt = parse_date_from_any(match.group(1))
        if end_dt:
            return end_dt - timedelta(days=6), end_dt

    return None, None


def extract_kpi_value(df: pd.DataFrame, key: str) -> Any:
    col = find_col(df, key)
    if not col:
        return None

    col_data = df[col]
    if isinstance(col_data, pd.DataFrame):
        col_data = col_data.iloc[:, 0]

    return first_non_empty(col_data)


def extract_client_name(
    sheets: Dict[str, pd.DataFrame],
    detected: Dict[str, pd.DataFrame],
    filename: str = "",
) -> Optional[str]:
    for table_name in ["campaign_kpi_summary", "campaign_delivery"]:
        df = detected.get(table_name)
        if df is not None:
            col = find_col(df, "client_brand")
            if col:
                val = first_non_empty(df[col])
                if clean_text(val):
                    return clean_text(val)

    for text in scan_workbook_texts(sheets):
        low = normalize_text(text)
        if "client:" in low or "brand:" in low:
            parts = text.split(":")
            if len(parts) > 1:
                candidate = clean_text(parts[-1])
                if candidate:
                    return candidate

    kpi_df = detected.get("campaign_kpi_summary")
    if kpi_df is not None:
        campaign_val = clean_text(extract_kpi_value(kpi_df, "campaign"))
        if campaign_val:
            for sep in [" - ", " | ", "_"]:
                if sep in campaign_val:
                    return clean_text(campaign_val.split(sep)[0])
            return campaign_val.split()[0] if campaign_val.split() else None

    if filename:
        base = Path(filename).stem
        if " - " in base:
            return clean_text(base.split(" - ")[0])

    return None


def join_dimension_values(df: pd.DataFrame, dim_key: str) -> Optional[str]:
    col = find_col(df, dim_key)
    if not col:
        return None

    col_data = df[col]
    if isinstance(col_data, pd.DataFrame):
        col_data = col_data.iloc[:, 0]

    if dim_key == "geo":
        raw_vals = []
        for v in col_data.tolist():
            cleaned = normalise_market_label(v)
            if cleaned:
                raw_vals.extend([x.strip() for x in cleaned.split(",") if x.strip()])
    else:
        raw_vals = [clean_dimension_label(v) for v in col_data.tolist() if clean_dimension_label(v)]

    blocked = {
        "total", "totals", "overall", "summary",
        "campaign", "format", "formats",
        "market", "markets"
    }

    raw_vals = [
        v for v in raw_vals
        if normalize_header(v) not in blocked
    ]

    if not raw_vals:
        return None

    base = raw_vals[0]
    cleaned = [base]

    for val in raw_vals[1:]:
        if dim_key == "geo":
            candidate = val
        else:
            parts_base = base.split(" - ")
            parts_val = val.split(" - ")

            suffix = parts_val

            for i in range(min(len(parts_base), len(parts_val))):
                if parts_base[i] != parts_val[i]:
                    suffix = parts_val[i:]
                    break
            else:
                if len(parts_val) > len(parts_base):
                    suffix = parts_val[len(parts_base):]
                else:
                    suffix = [val]

            candidate = clean_dimension_label(" - ".join([p for p in suffix if p]))

        if candidate:
            cleaned.append(candidate)

    seen = []
    final = []

    for v in cleaned:
        if v not in seen:
            seen.append(v)
            final.append(v)

    return ", ".join(final)


def extract_top_titles(df: pd.DataFrame, metric_key: str, max_rank: int = 5) -> Dict[str, Any]:
    site_col = find_col(df, "site")
    metric_col = find_col(df, metric_key)

    if not site_col or not metric_col:
        return {}

    site_data = df[site_col]
    if isinstance(site_data, pd.DataFrame):
        site_data = site_data.iloc[:, 0]

    metric_data = df[metric_col]
    if isinstance(metric_data, pd.DataFrame):
        metric_data = metric_data.iloc[:, 0]

    temp = pd.DataFrame({"site": site_data, "metric": metric_data})
    temp["site_clean"] = temp["site"].apply(clean_text)
    temp["site_norm"] = temp["site_clean"].apply(lambda x: normalize_header(x))
    temp["metric"] = temp["metric"].apply(safe_number)
    temp = temp.dropna(subset=["metric"])

    blocked_exact = {
        "site", "publisher", "domain", "environment", "property",
        "placement", "title", "inventory", "app", "website",
        "total", "totals", "overall", "average", "avg", "summary",
        "campaign", "campaign name", "market", "format", "creative",
    }

    temp = temp[~temp["site_norm"].isin(blocked_exact)]
    temp = temp[temp["site_clean"] != ""]

    if temp.empty:
        return {}

    temp = temp.sort_values(by="metric", ascending=False).head(max_rank)

    prefix = {
        "ctr": "TOP_TITLES_CTR",
        "engagement_rate": "TOP_TITLES_ER",
        "vcr": "TOP_TITLES_VCR",
    }[metric_key]

    out = {}

    for idx, (_, row) in enumerate(temp.iterrows(), start=1):
        out[f"{prefix}_{idx}_NAME"] = row["site_clean"]
        out[f"{prefix}_{idx}_VALUE"] = format_percent(row["metric"])

    return out


def extract_top_markets(df: pd.DataFrame, metric_key: str, max_rank: int = 5) -> Dict[str, Any]:
    geo_col = find_col(df, "geo")
    metric_col = find_col(df, metric_key)

    if not geo_col or not metric_col:
        return {}

    geo_data = df[geo_col]
    if isinstance(geo_data, pd.DataFrame):
        geo_data = geo_data.iloc[:, 0]

    metric_data = df[metric_col]
    if isinstance(metric_data, pd.DataFrame):
        metric_data = metric_data.iloc[:, 0]

    temp = pd.DataFrame({"market": geo_data, "metric": metric_data})
    temp["market_clean"] = temp["market"].apply(normalise_market_label)
    temp["market_norm"] = temp["market_clean"].apply(lambda x: normalize_header(x))
    temp["metric"] = temp["metric"].apply(safe_number)
    temp = temp.dropna(subset=["metric"])

    blocked_exact = {
        "geo", "market", "markets", "country", "countries", "region",
        "territory", "location", "locale", "total", "totals", "overall",
        "average", "avg", "summary", "campaign", "campaign name",
    }

    temp = temp[~temp["market_norm"].isin(blocked_exact)]
    temp = temp[temp["market_clean"] != ""]

    if temp.empty:
        return {}

    temp = temp.sort_values(by="metric", ascending=False).head(max_rank)

    prefix = {
        "ctr": "TOP_MARKETS_CTR",
        "engagement_rate": "TOP_MARKETS_ER",
        "vcr": "TOP_MARKETS_VCR",
    }[metric_key]

    out = {}

    for idx, (_, row) in enumerate(temp.iterrows(), start=1):
        out[f"{prefix}_{idx}_NAME"] = row["market_clean"]
        out[f"{prefix}_{idx}_VALUE"] = format_percent(row["metric"])

    return out



def limit_exec_summary(value: Optional[str], max_chars: int = EXEC_SUMMARY_MAX_CHARS) -> str:
    """Keep injected summaries short enough for the PPT text container."""
    text = clean_text(value)
    if not text:
        return MISSING_PPT_VALUE
    if len(text) <= max_chars:
        return text
    trimmed = text[:max_chars].rsplit(" ", 1)[0].strip(" ,.;:-")
    return trimmed + "..."


def value_available(value: Any) -> bool:
    text = clean_text(value)
    return bool(text and text not in {MISSING_PPT_VALUE, MISSING_DISPLAY_VALUE, "Not available", "N/A"})


def build_default_exec_summary(mapped: Dict[str, Any]) -> str:
    """Create a safe default summary when GPT has not supplied/selected one.

    Validation creates a default summary in chat, but the One Pager/Slide Deck
    action does not always receive it as request.exec_summary. This fallback
    prevents {{EXEC_SUMMARY}} or lorem ipsum from remaining in the PPT.
    """
    campaign = clean_text(mapped.get("CAMPAIGN_NAME")) or "The campaign"
    markets = clean_text(mapped.get("CAMPAIGN_MARKETS"))
    impressions = clean_text(mapped.get("DELIVERED_IMPRESSIONS"))
    ctr = clean_text(mapped.get("PERFORMANCE_CTR"))
    er = clean_text(mapped.get("PERFORMANCE_ENGAGEMENT_RATE"))
    vcr = clean_text(mapped.get("PERFORMANCE_VCR"))
    on_screen = clean_text(mapped.get("PERFORMANCE_ON_SCREEN"))

    parts = []

    opening = campaign
    if value_available(markets):
        opening += f" delivered across {markets}"
    if value_available(impressions):
        opening += f", generating {impressions} impressions"
    opening += "."
    parts.append(opening)

    metrics = []
    if value_available(ctr):
        metrics.append(f"CTR of {ctr}")
    if value_available(er):
        metrics.append(f"engagement rate of {er}")
    if value_available(vcr):
        metrics.append(f"VCR of {vcr}")
    if value_available(on_screen):
        metrics.append(f"on-screen rate of {on_screen}")

    if metrics:
        if len(metrics) == 1:
            metric_sentence = f"The campaign achieved a {metrics[0]}."
        else:
            metric_sentence = "The campaign achieved " + ", ".join(metrics[:-1]) + f" and {metrics[-1]}."
        parts.append(metric_sentence)

    top_ctr = clean_text(mapped.get("TOP_TITLES_CTR_1_NAME"))
    top_er = clean_text(mapped.get("TOP_TITLES_ER_1_NAME"))
    top_vcr = clean_text(mapped.get("TOP_TITLES_VCR_1_NAME"))

    leaders = []
    if value_available(top_ctr):
        leaders.append(f"{top_ctr} led CTR")
    if value_available(top_er) and top_er != top_ctr:
        leaders.append(f"{top_er} led engagement")
    if value_available(top_vcr) and top_vcr not in {top_ctr, top_er}:
        leaders.append(f"{top_vcr} led video completion")

    if leaders:
        parts.append(" ".join(leaders) + ".")

    return limit_exec_summary(" ".join(parts))


def resolve_exec_summary_for_ppt(request_summary: Optional[str], mapped: Dict[str, Any]) -> str:
    supplied = clean_text(request_summary)
    if supplied:
        return limit_exec_summary(supplied)
    return build_default_exec_summary(mapped)




def chat_display_value(value: Any) -> str:
    """Return a clean chat-facing value while keeping PPT N/A logic separate."""
    text = clean_text(value)
    if not text or text in {MISSING_PPT_VALUE, MISSING_DISPLAY_VALUE, "nan", "None"}:
        return "Not available"
    return text.replace("—", "-")


def chat_top_rows(mapped: Dict[str, Any], prefix: str, max_rows: int = 3) -> List[str]:
    rows = []
    for i in range(1, max_rows + 1):
        name = mapped.get(f"{prefix}_{i}_NAME")
        value = mapped.get(f"{prefix}_{i}_VALUE")
        if value_available(name) and value_available(value):
            rows.append(f"• {chat_display_value(name)} - {chat_display_value(value)}")
    return rows or ["• Not available"]


def valid_top_row_count(mapped: Dict[str, Any], prefix: str, max_rows: int = 5) -> int:
    count = 0
    for i in range(1, max_rows + 1):
        if value_available(mapped.get(f"{prefix}_{i}_NAME")) and value_available(mapped.get(f"{prefix}_{i}_VALUE")):
            count += 1
    return count


def has_minimum_market_data(mapped: Dict[str, Any], min_rows: int = 3) -> bool:
    return any([
        valid_top_row_count(mapped, "TOP_MARKETS_CTR", 5) >= min_rows,
        valid_top_row_count(mapped, "TOP_MARKETS_ER", 5) >= min_rows,
        valid_top_row_count(mapped, "TOP_MARKETS_VCR", 5) >= min_rows,
    ])


def render_modular_section_options(section_availability: Dict[str, Any]) -> str:
    labels_by_number = {item.get("number"): item.get("label") for item in section_availability.get("selection_labels", [])}
    section_names = {
        1: labels_by_number.get(1, "Title Performance"),
        2: labels_by_number.get(2, "Market Performance"),
        3: labels_by_number.get(3, "Top Titles & Markets"),
        4: labels_by_number.get(4, "Creative Overview"),
        5: labels_by_number.get(5, "Creative Performance"),
        6: "Attention Score",
        7: "Happydemics",
        8: "Lumen Results",
        9: "Brand Study",
        10: "Learning Recommendations",
    }
    lines = [
        "Available Modular Sections",
        "",
    ]
    for number in range(1, 11):
        lines.append(f"{number}. {section_names[number]}")
    return "\n".join(lines)


def build_summary_inputs(mapped: Dict[str, Any]) -> Dict[str, Any]:
    """Return a compact, GPT-friendly set of values for writing the Exec Summary.

    The API owns validation rendering, but GPT should own the copywriting.
    This compact payload gives GPT enough validated facts to write a polished
    summary without returning the full mapped/display dictionaries.
    """
    def top(prefix: str, max_rows: int = 3) -> List[Dict[str, str]]:
        rows = []
        for i in range(1, max_rows + 1):
            name = mapped.get(f"{prefix}_{i}_NAME")
            value = mapped.get(f"{prefix}_{i}_VALUE")
            if value_available(name) and value_available(value):
                rows.append({
                    "name": chat_display_value(name),
                    "value": chat_display_value(value),
                })
        return rows

    return {
        "campaign_name": chat_display_value(mapped.get("CAMPAIGN_NAME")),
        "client": chat_display_value(mapped.get("CLIENT_NAME") or mapped.get("CLIENT")),
        "markets": chat_display_value(mapped.get("CAMPAIGN_MARKETS")),
        "live_dates": chat_display_value(mapped.get("LIVE_DATES_FULL")),
        "campaign_period": chat_display_value(mapped.get("CAMPAIGN_PERIOD")),
        "delivered_impressions": chat_display_value(mapped.get("DELIVERED_IMPRESSIONS")),
        "io_overall_impressions": chat_display_value(mapped.get("IO_OVERALL_IMPRESSIONS")),
        "added_value_impressions": chat_display_value(mapped.get("ADDED_VALUE_IMPRESSIONS")),
        "delivery_incl_av": chat_display_value(mapped.get("DELIVERY_WITH_AV_PERCENT")),
        "added_value_worth": chat_display_value(mapped.get("ADDED_VALUE_WORTH")),
        "budget": chat_display_value(mapped.get("CAMPAIGN_BUDGET")),
        "ctr": chat_display_value(mapped.get("PERFORMANCE_CTR")),
        "engagement_rate": chat_display_value(mapped.get("PERFORMANCE_ENGAGEMENT_RATE")),
        "vcr": chat_display_value(mapped.get("PERFORMANCE_VCR")),
        "on_screen_rate": chat_display_value(mapped.get("PERFORMANCE_ON_SCREEN")),
        "creative_formats": chat_display_value(mapped.get("CAMPAIGN_FORMATS")),
        "top_ctr": top("TOP_TITLES_CTR", 3),
        "top_engagement_rate": top("TOP_TITLES_ER", 3),
        "top_vcr": top("TOP_TITLES_VCR", 3),
        "top_market_ctr": top("TOP_MARKETS_CTR", 3),
        "top_market_engagement_rate": top("TOP_MARKETS_ER", 3),
        "top_market_vcr": top("TOP_MARKETS_VCR", 3),
        "market_data_available": has_minimum_market_data(mapped, 3),
    }


def render_validation_text(mapped: Dict[str, Any]) -> str:
    """Render validation sections only.

    Exec Summary is intentionally excluded so GPT can write it with a better
    client-ready tone, then store/pass it as the active exec_summary for PPTs.
    """
    lines = [
        "**Campaign Overview**",
        "",
        f"• **Campaign Name:** {chat_display_value(mapped.get('CAMPAIGN_NAME'))}",
        f"• **Client:** {chat_display_value(mapped.get('CLIENT_NAME') or mapped.get('CLIENT'))}",
        f"• **Markets:** {chat_display_value(mapped.get('CAMPAIGN_MARKETS'))}",
        f"• **Live Dates:** {chat_display_value(mapped.get('LIVE_DATES_FULL'))}",
        f"• **Campaign Period:** {chat_display_value(mapped.get('CAMPAIGN_PERIOD'))}",
        "",
        "**Delivery & Spend**",
        "",
        f"• **Delivered Impressions:** {chat_display_value(mapped.get('DELIVERED_IMPRESSIONS'))}",
        f"• **IO Overall Impressions:** {chat_display_value(mapped.get('IO_OVERALL_IMPRESSIONS'))}",
        f"• **Added Value Impressions:** {chat_display_value(mapped.get('ADDED_VALUE_IMPRESSIONS'))}",
        f"• **Delivery (incl. AV):** {chat_display_value(mapped.get('DELIVERY_WITH_AV_PERCENT'))}",
        f"• **Added Value Worth:** {chat_display_value(mapped.get('ADDED_VALUE_WORTH'))}",
        f"• **Budget:** {chat_display_value(mapped.get('CAMPAIGN_BUDGET'))}",
        "",
        "**Performance Metrics**",
        "",
        f"• **CTR:** {chat_display_value(mapped.get('PERFORMANCE_CTR'))}",
        f"• **Engagement Rate:** {chat_display_value(mapped.get('PERFORMANCE_ENGAGEMENT_RATE'))}",
        f"• **VCR:** {chat_display_value(mapped.get('PERFORMANCE_VCR'))}",
        f"• **On-Screen Rate:** {chat_display_value(mapped.get('PERFORMANCE_ON_SCREEN'))}",
        "",
        "**Creative Formats**",
        "",
        f"• {chat_display_value(mapped.get('CAMPAIGN_FORMATS'))}",
        "",
        "**Top Performing Titles**",
        "",
        "CTR",
        *chat_top_rows(mapped, "TOP_TITLES_CTR", 3),
        "",
        "ER",
        *chat_top_rows(mapped, "TOP_TITLES_ER", 3),
        "",
        "VCR",
        *chat_top_rows(mapped, "TOP_TITLES_VCR", 3),
    ]

    if has_minimum_market_data(mapped, 3):
        lines.extend([
            "",
            "**Top Performing Markets**",
            "",
            "CTR",
            *chat_top_rows(mapped, "TOP_MARKETS_CTR", 3),
            "",
            "ER",
            *chat_top_rows(mapped, "TOP_MARKETS_ER", 3),
            "",
            "VCR",
            *chat_top_rows(mapped, "TOP_MARKETS_VCR", 3),
        ])

    return "\n".join(lines)

def copy_value(mapped: Dict[str, Any], source_key: str) -> str:
    value = mapped.get(source_key)
    if value is None or clean_text(value) == "":
        return MISSING_PPT_VALUE
    return str(value)


def add_placeholder_aliases(mapped: Dict[str, Any]) -> Dict[str, Any]:
    """Add aliases used by combined Top Titles & Markets sections.

    Some templates use a dedicated placeholder set for the combined section rather
    than the standalone TOP_TITLES_* and TOP_MARKETS_* keys. This mirrors top
    publisher/title values into the combined placeholder names and guarantees
    unavailable market rows resolve to N/A instead of template defaults such as
    Market 1 or XX%.
    """
    metric_map = {
        "CTR": "CTR",
        "ER": "ER",
        "VCR": "VCR",
    }

    for metric in metric_map:
        title_prefix = f"TOP_TITLES_{metric}"
        market_prefix = f"TOP_MARKETS_{metric}"

        for i in range(1, 6):
            title_name = copy_value(mapped, f"{title_prefix}_{i}_NAME")
            title_value = copy_value(mapped, f"{title_prefix}_{i}_VALUE")
            market_name = copy_value(mapped, f"{market_prefix}_{i}_NAME")
            market_value = copy_value(mapped, f"{market_prefix}_{i}_VALUE")

            title_alias_prefixes = [
                f"TOP_TITLES_MARKETS_{metric}_{i}",
                f"TOP_TITLES_AND_MARKETS_{metric}_{i}",
                f"TOP_TITLES_MARKETS_TITLE_{metric}_{i}",
                f"TOP_TITLES_MARKETS_TITLES_{metric}_{i}",
                f"TOP_TITLES_MARKETS_PERFORMER_{metric}_{i}",
                f"TTM_TITLE_{metric}_{i}",
                f"TTM_TITLES_{metric}_{i}",
            ]

            market_alias_prefixes = [
                f"TOP_TITLES_MARKETS_MARKET_{metric}_{i}",
                f"TOP_TITLES_MARKETS_MARKETS_{metric}_{i}",
                f"TOP_TITLES_AND_MARKETS_MARKET_{metric}_{i}",
                f"TOP_TITLES_AND_MARKETS_MARKETS_{metric}_{i}",
                f"TTM_MARKET_{metric}_{i}",
                f"TTM_MARKETS_{metric}_{i}",
            ]

            for alias_prefix in title_alias_prefixes:
                mapped.setdefault(f"{alias_prefix}_NAME", title_name)
                mapped.setdefault(f"{alias_prefix}_VALUE", title_value)

            for alias_prefix in market_alias_prefixes:
                mapped.setdefault(f"{alias_prefix}_NAME", market_name)
                mapped.setdefault(f"{alias_prefix}_VALUE", market_value)

            # Common shorter row placeholders sometimes used inside combined blocks.
            mapped.setdefault(f"TITLE_{metric}_{i}_NAME", title_name)
            mapped.setdefault(f"TITLE_{metric}_{i}_VALUE", title_value)
            mapped.setdefault(f"MARKET_{metric}_{i}_NAME", market_name)
            mapped.setdefault(f"MARKET_{metric}_{i}_VALUE", market_value)

            # Conservative extra aliases for templates that use metric-first naming.
            # Kept deliberately small to avoid large Action responses.
            mapped.setdefault(f"{metric}_TITLE_{i}_NAME", title_name)
            mapped.setdefault(f"{metric}_TITLE_{i}_VALUE", title_value)
            mapped.setdefault(f"{metric}_MARKET_{i}_NAME", market_name)
            mapped.setdefault(f"{metric}_MARKET_{i}_VALUE", market_value)
            mapped.setdefault(f"TOP_{metric}_TITLE_{i}_NAME", title_name)
            mapped.setdefault(f"TOP_{metric}_TITLE_{i}_VALUE", title_value)
            mapped.setdefault(f"TOP_{metric}_MARKET_{i}_NAME", market_name)
            mapped.setdefault(f"TOP_{metric}_MARKET_{i}_VALUE", market_value)

    return mapped


def replace_template_default_literals(text: str) -> str:
    """Replace visible template defaults that are not written as {{placeholders}}."""
    replacements = {
        "Title 1": MISSING_PPT_VALUE,
        "Title 2": MISSING_PPT_VALUE,
        "Title 3": MISSING_PPT_VALUE,
        "Title 4": MISSING_PPT_VALUE,
        "Title 5": MISSING_PPT_VALUE,
        "Market 1": MISSING_PPT_VALUE,
        "Market 2": MISSING_PPT_VALUE,
        "Market 3": MISSING_PPT_VALUE,
        "Market 4": MISSING_PPT_VALUE,
        "Market 5": MISSING_PPT_VALUE,
        "XX%": MISSING_PPT_VALUE,
        "XX": MISSING_PPT_VALUE,
    }
    new_text = text
    for old, new in replacements.items():
        if old in new_text:
            new_text = new_text.replace(old, new)
    return new_text

def validate_mapped_values(mapped: Dict[str, Any]) -> Dict[str, Any]:
    required_core = [
        "CAMPAIGN_NAME",
        "DELIVERED_IMPRESSIONS",
        "PERFORMANCE_CTR",
        "PERFORMANCE_ENGAGEMENT_RATE",
        "PERFORMANCE_VCR",
        "LIVE_DATES_FULL",
    ]

    missing_required = [
        k for k in required_core
        if not mapped.get(k) or mapped.get(k) == MISSING_PPT_VALUE
    ]

    warnings = []
    for key, label in [
        ("CLIENT_NAME", "CLIENT_NAME missing"),
        ("CAMPAIGN_MARKETS", "CAMPAIGN_MARKETS missing"),
        ("CAMPAIGN_FORMATS", "CAMPAIGN_FORMATS missing"),
        ("TOP_TITLES_CTR_1_NAME", "Top Titles CTR missing"),
        ("TOP_TITLES_VCR_1_NAME", "Top Titles VCR missing"),
        ("TOP_TITLES_ER_1_NAME", "Top Titles ER missing"),
        ("TOP_MARKETS_CTR_1_NAME", "Top Markets CTR missing"),
        ("TOP_MARKETS_VCR_1_NAME", "Top Markets VCR missing"),
        ("TOP_MARKETS_ER_1_NAME", "Top Markets ER missing"),
    ]:
        if mapped.get(key) == MISSING_PPT_VALUE:
            warnings.append(label)

    return {
        "is_valid": len(missing_required) == 0,
        "missing_required": missing_required,
        "warnings": warnings,
    }


def has_ppt_value(mapped: Dict[str, Any], key: str) -> bool:
    value = mapped.get(key)
    return value is not None and clean_text(value) not in ["", MISSING_PPT_VALUE, MISSING_DISPLAY_VALUE]


def build_section_availability(
    mapped: Dict[str, Any],
    detected: Dict[str, pd.DataFrame],
) -> Dict[str, Any]:
    """Return data availability for data-driven modular sections only.

    Manual/add-on sections are intentionally excluded:
    - ATTENTION_SCORE
    - HAPPYDEMICS
    - LUMEN_RESULTS
    - BRAND_STUDY
    - LEARNING_RECOMMENDATIONS
    """

    has_campaign_kpis = all([
        has_ppt_value(mapped, "CAMPAIGN_NAME"),
        has_ppt_value(mapped, "DELIVERED_IMPRESSIONS"),
        has_ppt_value(mapped, "PERFORMANCE_CTR"),
        has_ppt_value(mapped, "PERFORMANCE_ENGAGEMENT_RATE"),
        has_ppt_value(mapped, "PERFORMANCE_VCR"),
    ])

    has_site_rankings = any([
        has_ppt_value(mapped, "TOP_TITLES_CTR_1_NAME"),
        has_ppt_value(mapped, "TOP_TITLES_ER_1_NAME"),
        has_ppt_value(mapped, "TOP_TITLES_VCR_1_NAME"),
    ])

    has_market_rankings = any([
        has_ppt_value(mapped, "TOP_MARKETS_CTR_1_NAME"),
        has_ppt_value(mapped, "TOP_MARKETS_ER_1_NAME"),
        has_ppt_value(mapped, "TOP_MARKETS_VCR_1_NAME"),
    ])

    # Market-based modules require 3+ valid market rows. A single campaign market
    # such as "FR" is useful context, but not enough for market performance sections.
    has_market_data = has_minimum_market_data(mapped, 3)

    format_df = detected.get("format")
    has_creative_data = has_ppt_value(mapped, "CAMPAIGN_FORMATS") or format_df is not None

    has_creative_performance_data = False
    if format_df is not None:
        has_creative_performance_data = any([
            find_col(format_df, "ctr") is not None,
            find_col(format_df, "engagement_rate") is not None,
            find_col(format_df, "vcr") is not None,
            find_col(format_df, "impressions") is not None,
        ])

    availability = {
        "TITLE_PERFORMANCE": {
            "label": "Title Performance",
            "available": has_campaign_kpis,
            "note": "" if has_campaign_kpis else "No Data Available",
        },
        "MARKET_PERFORMANCE": {
            "label": "Market Performance",
            "available": has_market_data,
            "note": "" if has_market_data else "No Data Available",
        },
        "TOP_TITLES_MARKETS": {
            "label": "Top Titles & Markets",
            "available": has_site_rankings and has_market_data,
            "note": "" if (has_site_rankings and has_market_data) else "No Data Available",
        },
        "CREATIVE_OVERVIEW": {
            "label": "Creative Overview",
            "available": has_creative_data,
            "note": "" if has_creative_data else "No Data Available",
        },
        "CREATIVE_PERFORMANCE": {
            "label": "Creative Performance",
            "available": has_creative_performance_data,
            "note": "" if has_creative_performance_data else "No Data Available",
        },
    }

    selection_labels = []
    ordered = [
        (1, "TITLE_PERFORMANCE"),
        (2, "MARKET_PERFORMANCE"),
        (3, "TOP_TITLES_MARKETS"),
        (4, "CREATIVE_OVERVIEW"),
        (5, "CREATIVE_PERFORMANCE"),
    ]

    for number, section_id in ordered:
        item = availability[section_id]
        label = item["label"]
        if item["note"]:
            label = f"{label} ({item['note']})"
        selection_labels.append({
            "number": number,
            "section_id": section_id,
            "label": label,
            "available": item["available"],
            "note": item["note"],
        })

    return {
        "sections": availability,
        "selection_labels": selection_labels,
        "notes": {
            "manual_sections": [
                "ATTENTION_SCORE",
                "HAPPYDEMICS",
                "LUMEN_RESULTS",
                "BRAND_STUDY",
                "LEARNING_RECOMMENDATIONS",
            ],
            "rule": "Only sections 1-5 are data-checked. Manual/add-on sections are not marked as unavailable.",
        },
    }


def load_rules_master() -> Optional[Dict[str, pd.DataFrame]]:
    if not RULES_MASTER_PATH.exists():
        return None

    try:
        xls = pd.ExcelFile(RULES_MASTER_PATH, engine="openpyxl")
        return {
            sheet: pd.read_excel(RULES_MASTER_PATH, sheet_name=sheet, engine="openpyxl")
            for sheet in xls.sheet_names
        }
    except Exception:
        return None


def build_mapped_values(sheets: Dict[str, pd.DataFrame], filename: str = "") -> Dict[str, Any]:
    rules = load_rules_master()
    detected = detect_tables(sheets)

    kpi_df = detected.get("campaign_kpi_summary")
    delivery_df = detected.get("campaign_delivery")
    site_df = detected.get("site")
    geo_df = detected.get("geo")
    format_df = detected.get("format")

    start_dt, end_dt = extract_dates(sheets, detected, filename)
    client_name = extract_client_name(sheets, detected, filename)

    mapped: Dict[str, Any] = {}

    if kpi_df is not None:
        mapped["CAMPAIGN_NAME"] = clean_text(extract_kpi_value(kpi_df, "campaign"))
        mapped["DELIVERED_IMPRESSIONS"] = format_number(safe_number(extract_kpi_value(kpi_df, "impressions")), 0)
        mapped["PERFORMANCE_CTR"] = format_percent(extract_kpi_value(kpi_df, "ctr"))
        mapped["PERFORMANCE_ENGAGEMENT_RATE"] = format_percent(extract_kpi_value(kpi_df, "engagement_rate"))
        mapped["PERFORMANCE_VCR"] = format_percent(extract_kpi_value(kpi_df, "vcr"))

        on_screen_val = None
        for key in ["mobkoi_on_screen", "mrc_viewability", "on_screen"]:
            col = find_col(kpi_df, key)
            if col:
                col_data = kpi_df[col]
                if isinstance(col_data, pd.DataFrame):
                    col_data = col_data.iloc[:, 0]
                on_screen_val = first_non_empty(col_data)
                break

        mapped["PERFORMANCE_ON_SCREEN"] = format_percent(on_screen_val)
        mapped["CAMPAIGN_BUDGET"] = format_currency(safe_number(extract_kpi_value(kpi_df, "spend")))
    else:
        mapped["CAMPAIGN_NAME"] = None
        mapped["DELIVERED_IMPRESSIONS"] = None
        mapped["PERFORMANCE_CTR"] = None
        mapped["PERFORMANCE_ENGAGEMENT_RATE"] = None
        mapped["PERFORMANCE_VCR"] = None
        mapped["PERFORMANCE_ON_SCREEN"] = None
        mapped["CAMPAIGN_BUDGET"] = None

    if delivery_df is not None:
        io_val = safe_number(extract_kpi_value(delivery_df, "sold_paid_units"))
        av_units_val = abs_number(extract_kpi_value(delivery_df, "delivered_overall_av_units"))
        av_amount_val = abs_number(extract_kpi_value(delivery_df, "delivered_av_amount"))

        mapped["IO_OVERALL_IMPRESSIONS"] = format_number(io_val, 0)
        mapped["DELIVERED_OVERALL_AV_UNITS"] = format_number(av_units_val, 0)
        mapped["DELIVERY_WITH_AV_PERCENT"] = format_percent(extract_kpi_value(delivery_df, "delivery_incl_av"))
        mapped["ADDED_VALUE_WORTH"] = format_currency(av_amount_val)
        mapped["ADDED_VALUE_IMPRESSIONS"] = format_number(av_units_val, 0)
    else:
        mapped["IO_OVERALL_IMPRESSIONS"] = None
        mapped["DELIVERED_OVERALL_AV_UNITS"] = None
        mapped["DELIVERY_WITH_AV_PERCENT"] = None
        mapped["ADDED_VALUE_WORTH"] = None
        mapped["ADDED_VALUE_IMPRESSIONS"] = None

    mapped["CLIENT_NAME"] = client_name
    mapped["CLIENT"] = client_name
    mapped["CAMPAIGN_FORMATS"] = join_dimension_values(format_df, "format") if format_df is not None else None
    mapped["CAMPAIGN_MARKETS"] = join_dimension_values(geo_df, "geo") if geo_df is not None else None
    mapped["LIVE_DATES_SHORT"] = format_date_short(start_dt, end_dt)
    mapped["LIVE_DATES_FULL"] = format_date_full(start_dt, end_dt)
    mapped["CAMPAIGN_PERIOD"] = format_quarter_from_date(end_dt)

    if site_df is not None:
        mapped.update(extract_top_titles(site_df, "ctr", 5))
        mapped.update(extract_top_titles(site_df, "engagement_rate", 5))
        mapped.update(extract_top_titles(site_df, "vcr", 5))

    if geo_df is not None:
        mapped.update(extract_top_markets(geo_df, "ctr", 5))
        mapped.update(extract_top_markets(geo_df, "engagement_rate", 5))
        mapped.update(extract_top_markets(geo_df, "vcr", 5))

    for metric_prefix in ["TOP_TITLES_CTR", "TOP_TITLES_ER", "TOP_TITLES_VCR"]:
        for i in range(1, 6):
            mapped.setdefault(f"{metric_prefix}_{i}_NAME", None)
            mapped.setdefault(f"{metric_prefix}_{i}_VALUE", None)

    for metric_prefix in ["TOP_MARKETS_CTR", "TOP_MARKETS_ER", "TOP_MARKETS_VCR"]:
        for i in range(1, 6):
            mapped.setdefault(f"{metric_prefix}_{i}_NAME", None)
            mapped.setdefault(f"{metric_prefix}_{i}_VALUE", None)

    mapped = add_placeholder_aliases(mapped)
    mapped = fill_missing_for_ppt(mapped)
    default_exec_summary = build_default_exec_summary(mapped)
    mapped["EXEC_SUMMARY"] = default_exec_summary
    display_values = build_display_values(mapped)
    validation = validate_mapped_values(mapped)
    section_availability = build_section_availability(mapped, detected)
    display_validation_text = render_validation_text(mapped)
    display_section_options = render_modular_section_options(section_availability)
    summary_inputs = build_summary_inputs(mapped)

    diagnostics = {
        "detected_tables": list(detected.keys()),
        "rules_master_loaded": rules is not None,
        "version": APP_VERSION,
    }

    return {
        "mapped_values": mapped,
        "display_values": display_values,
        "display_validation_text": display_validation_text,
        "default_exec_summary": default_exec_summary,
        "summary_inputs": summary_inputs,
        "validation": validation,
        "section_availability": section_availability,
        "display_section_options": display_section_options,
        "diagnostics": diagnostics,
    }


def compact_parsed_result(parsed: Dict[str, Any]) -> Dict[str, Any]:
    """Return a compact validation payload for ChatGPT Actions.

    The full mapped/display dictionaries are useful internally for PPT generation,
    but returning them through Actions can create oversized responses and unstable
    upload retries. GPT only needs the rendered validation text, default summary,
    validation flags, and section availability.
    """
    diagnostics = parsed.get("diagnostics") or {}
    compact_diagnostics = {
        "detected_tables": diagnostics.get("detected_tables", []),
        "version": diagnostics.get("version", APP_VERSION),
    }

    return {
        "display_validation_text": parsed.get("display_validation_text", ""),
        "default_exec_summary": parsed.get("default_exec_summary", ""),
        "summary_inputs": parsed.get("summary_inputs", {}),
        "validation": parsed.get("validation", {}),
        "section_availability": parsed.get("section_availability", {}),
        "display_section_options": parsed.get("display_section_options", ""),
        "diagnostics": compact_diagnostics,
    }


async def download_openai_file(file_ref: OpenAIFileRef, dest_dir: Path) -> Path:
    """Download an OpenAI Actions file reference with quiet internal retries.

    This reduces visible ChatGPT Actions errors caused by temporary file-link readiness,
    Render cold starts, or transient signed URL fetch failures. The function only raises
    after all server-side retries have been exhausted.
    """
    if not file_ref or not file_ref.download_link:
        raise HTTPException(status_code=400, detail="Missing download_link in openaiFileIdRefs.")

    safe_name = file_ref.name or "uploaded_eoc.xlsx"
    safe_name = re.sub(r"[^A-Za-z0-9._ -]+", "_", safe_name)
    output_path = dest_dir / safe_name

    last_error = None
    retry_delays = [0, 1.5, 3.0, 5.0]

    async with httpx.AsyncClient(timeout=90.0, follow_redirects=True) as client:
        for attempt, delay in enumerate(retry_delays, start=1):
            if delay:
                await asyncio.sleep(delay)

            try:
                response = await client.get(file_ref.download_link)
                response.raise_for_status()

                if not response.content:
                    raise ValueError("Downloaded file was empty.")

                output_path.write_bytes(response.content)
                return output_path

            except Exception as exc:
                last_error = exc
                continue

    raise HTTPException(
        status_code=400,
        detail=f"Could not download uploaded EOC file after internal retries: {last_error}"
    )


def read_uploaded_eoc(path: Path) -> Dict[str, pd.DataFrame]:
    try:
        return pd.read_excel(path, sheet_name=None, header=None, engine="openpyxl")
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Could not read EOC Excel file: {exc}")


def load_section_registry() -> Dict[str, Any]:
    if not SECTION_REGISTRY_PATH.exists():
        raise HTTPException(status_code=500, detail=f"Section registry not found at {SECTION_REGISTRY_PATH}")

    return json.loads(SECTION_REGISTRY_PATH.read_text(encoding="utf-8"))


def get_shape_text(shape) -> str:
    if not hasattr(shape, "text"):
        return ""
    return shape.text or ""


def is_section_id_shape(shape) -> bool:
    return get_shape_text(shape).strip().startswith("SECTION_ID:")


def extract_section_id_from_slide(slide) -> Optional[str]:
    for shape in slide.shapes:
        text = get_shape_text(shape).strip()
        if text.startswith("SECTION_ID:"):
            return text.replace("SECTION_ID:", "").strip()
    return None


def detect_template_sections(prs: Presentation) -> Dict[str, int]:
    detected = {}

    for index, slide in enumerate(prs.slides):
        section_id = extract_section_id_from_slide(slide)
        if section_id:
            detected[section_id] = index

    return detected


def ordered_selected_sections(registry: Dict[str, Any], selected_sections: List[str]) -> List[str]:
    selected_clean = [s.strip().upper() for s in selected_sections if s.strip()]
    final_sections = []

    for section_id, meta in registry.items():
        if meta.get("required") is True:
            final_sections.append(section_id)

    for section_id in selected_clean:
        if section_id not in final_sections:
            final_sections.append(section_id)

    final_sections = [section_id for section_id in final_sections if section_id in registry]
    final_sections.sort(key=lambda section_id: registry[section_id].get("default_order", 999))

    return final_sections


def find_main_group_shape(slide):
    candidates = []

    for shape in slide.shapes:
        if is_section_id_shape(shape):
            continue

        is_group = shape.shape_type == 6
        area = int(shape.width) * int(shape.height)

        candidates.append({
            "shape": shape,
            "is_group": is_group,
            "area": area
        })

    if not candidates:
        return None

    group_candidates = [c for c in candidates if c["is_group"]]

    if group_candidates:
        return max(group_candidates, key=lambda c: c["area"])["shape"]

    return max(candidates, key=lambda c: c["area"])["shape"]


def normalise_placeholder_key(value: Any) -> str:
    """Normalise PPT placeholder keys for robust lookup.

    PowerPoint can insert line breaks or split text runs inside tokens, for example
    {{TOP_TITLES_CTR_1_NAM
E}}. The template still visually represents the same
    placeholder, but a direct string replace will fail. This normalises the token
    content so split/line-wrapped placeholders resolve correctly.
    """
    key = str(value or "")
    key = key.strip().strip("{}").upper()
    key = re.sub(r"[^A-Z0-9_]+", "", key)
    return key


def build_placeholder_lookup(placeholder_values: Dict[str, Any]) -> Dict[str, str]:
    lookup: Dict[str, str] = {}

    for key, value in (placeholder_values or {}).items():
        if str(key).startswith("__"):
            continue

        norm_key = normalise_placeholder_key(key)
        if not norm_key:
            continue

        replacement = MISSING_PPT_VALUE if value is None or clean_text(value) == "" else str(value)
        lookup[norm_key] = replacement

        # PowerPoint sometimes clips or splits text-box placeholders in narrow
        # table cells, especially inside the combined Top Titles & Markets block.
        # In those templates, a visible token can resolve as *_NAM or *_VALU
        # instead of *_NAME / *_VALUE. Mirror each populated key to those safe
        # aliases so the combined section does not fall back to N/A while the
        # standalone sections still work normally.
        if norm_key.endswith("_NAME"):
            lookup.setdefault(norm_key[:-1], replacement)  # *_NAM
        if norm_key.endswith("_VALUE"):
            lookup.setdefault(norm_key[:-1], replacement)  # *_VALU

    return lookup


def replace_placeholders_in_text(text: str, placeholder_values: Dict[str, Any]) -> str:
    """Replace placeholders in a full text string, including split tokens.

    This handles both normal tokens like {{TOP_TITLES_CTR_1_NAME}} and tokens
    where PowerPoint has inserted spaces/newlines inside the braces. Unknown
    placeholders are converted to N/A as a final safety net.
    """
    if not text:
        return text

    lookup = build_placeholder_lookup(placeholder_values)

    def repl(match):
        raw_key = match.group(1)
        norm_key = normalise_placeholder_key(raw_key)
        return lookup.get(norm_key, MISSING_PPT_VALUE)

    # DOTALL allows placeholder text broken across lines/runs to be matched once
    # the paragraph text has been joined together.
    new_text = re.sub(r"\{\{(.*?)\}\}", repl, text, flags=re.DOTALL)

    # Direct replacement fallback for any simple placeholders missed above.
    for norm_key, replacement in lookup.items():
        new_text = new_text.replace("{{" + norm_key + "}}", replacement)

    new_text = replace_template_default_literals(new_text)
    return new_text


def replace_text_in_shape(shape, placeholder_values: Dict[str, Any]) -> None:
    if hasattr(shape, "text_frame"):
        for paragraph in shape.text_frame.paragraphs:
            runs = list(paragraph.runs)
            if not runs:
                continue

            # First try replacing across the whole paragraph. This handles
            # placeholders split across multiple PowerPoint text runs.
            original_full_text = "".join(run.text or "" for run in runs)
            replaced_full_text = replace_placeholders_in_text(original_full_text, placeholder_values)

            if replaced_full_text != original_full_text:
                runs[0].text = replaced_full_text
                for run in runs[1:]:
                    run.text = ""
                continue

            # Fallback for simple run-contained placeholders.
            for run in runs:
                if not run.text:
                    continue
                run.text = replace_placeholders_in_text(run.text, placeholder_values)

    if hasattr(shape, "shapes"):
        for subshape in shape.shapes:
            replace_text_in_shape(subshape, placeholder_values)


def replace_placeholders_on_slide(slide, placeholder_values: Dict[str, Any]) -> None:
    if not placeholder_values:
        return

    for shape in slide.shapes:
        replace_text_in_shape(shape, placeholder_values)


def create_output_presentation_with_source_theme(source_template_path: Path, slide_height: int) -> Presentation:
    output_prs = Presentation(str(source_template_path))
    output_prs.slide_height = slide_height

    slide_id_list = output_prs.slides._sldIdLst

    for slide_id in list(slide_id_list):
        output_prs.part.drop_rel(slide_id.rId)
        slide_id_list.remove(slide_id)

    return output_prs


def collect_relationship_ids_from_element(element) -> List[str]:
    rel_ids = set()

    for node in element.iter():
        for attr_name, attr_value in node.attrib.items():
            if attr_name.startswith("{" + REL_NS + "}"):
                rel_ids.add(attr_value)

    return list(rel_ids)


def remap_relationship_ids_in_element(element, rel_id_map: Dict[str, str]) -> None:
    for node in element.iter():
        for attr_name, attr_value in list(node.attrib.items()):
            if attr_value in rel_id_map:
                node.attrib[attr_name] = rel_id_map[attr_value]


def copy_relationships_for_element(source_slide, target_slide, copied_element) -> Dict[str, str]:
    rel_id_map = {}
    rel_ids = collect_relationship_ids_from_element(copied_element)

    for old_rid in rel_ids:
        try:
            source_rel = source_slide.part.rels[old_rid]
        except KeyError:
            continue

        try:
            if getattr(source_rel, "is_external", False):
                new_rid = target_slide.part.relate_to(
                    source_rel.target_ref,
                    source_rel.reltype,
                    is_external=True
                )
            else:
                new_rid = target_slide.part.relate_to(
                    source_rel.target_part,
                    source_rel.reltype
                )

            rel_id_map[old_rid] = new_rid

        except Exception:
            continue

    remap_relationship_ids_in_element(copied_element, rel_id_map)

    return rel_id_map


def copy_group_to_slide_relationship_safe(source_shape, source_slide, target_slide, new_left: int, new_top: int):
    copied_element = deepcopy(source_shape.element)

    copy_relationships_for_element(
        source_slide=source_slide,
        target_slide=target_slide,
        copied_element=copied_element
    )

    target_slide.shapes._spTree.insert_element_before(copied_element, "p:extLst")

    copied_shape = list(target_slide.shapes)[-1]

    try:
        copied_shape.left = new_left
        copied_shape.top = new_top
    except Exception:
        pass

    return copied_shape


def px_to_emu(px: int) -> int:
    return int(px * EMU_PER_PX)


def calculate_section_left(section_id: str, section_width: int, slide_width: int) -> int:
    if section_id == "TITLE_OVERVIEW":
        return int(slide_width - section_width)

    return int((slide_width - section_width) / 2)


def safe_filename(value: str) -> str:
    value = clean_text(value) or "Campaign"
    return re.sub(r'[\\/*?:"<>|]', "", value)


def build_grouped_stacked_modular_ppt(
    selected_sections: List[str],
    placeholder_values: Optional[Dict[str, Any]] = None,
    top_margin_px: int = 100,
    bottom_margin_px: int = 100,
    section_spacing_px: int = 60,
    template_path: Optional[Path] = None,
) -> Dict[str, Any]:

    tpl = template_path or MODULAR_TEMPLATE_PATH
    if not tpl.exists():
        raise HTTPException(status_code=500, detail=f"Modular template not found at {tpl}")

    registry = load_section_registry()
    source_prs = Presentation(str(tpl))
    detected_sections = detect_template_sections(source_prs)
    sections_to_build = ordered_selected_sections(registry, selected_sections)

    missing_sections = [
        section_id for section_id in sections_to_build
        if section_id not in detected_sections
    ]

    if missing_sections:
        raise HTTPException(
            status_code=400,
            detail={
                "message": "Some requested sections were not found in modular_sections_master.pptx",
                "missing_sections": missing_sections,
                "detected_sections": list(detected_sections.keys())
            }
        )

    section_data = []

    for section_id in sections_to_build:
        source_slide = source_prs.slides[detected_sections[section_id]]
        group_shape = find_main_group_shape(source_slide)

        if group_shape is None:
            continue

        section_data.append({
            "section_id": section_id,
            "label": registry[section_id].get("label", section_id),
            "slide": source_slide,
            "shape": group_shape,
            "width": group_shape.width,
            "height": group_shape.height
        })

    top_margin = px_to_emu(top_margin_px)
    bottom_margin = px_to_emu(bottom_margin_px)
    spacing = px_to_emu(section_spacing_px)

    total_height = top_margin + bottom_margin

    for i, section in enumerate(section_data):
        total_height += section["height"]
        if i < len(section_data) - 1:
            total_height += spacing

    output_prs = create_output_presentation_with_source_theme(
        source_template_path=tpl,
        slide_height=total_height
    )

    slide_width = output_prs.slide_width
    blank_layout = output_prs.slide_layouts[6]
    output_slide = output_prs.slides.add_slide(blank_layout)

    cursor_y = top_margin
    built_sections = []

    for section in section_data:
        section_id = section["section_id"]

        section_left = calculate_section_left(
            section_id=section_id,
            section_width=section["width"],
            slide_width=slide_width
        )

        copy_group_to_slide_relationship_safe(
            source_shape=section["shape"],
            source_slide=section["slide"],
            target_slide=output_slide,
            new_left=section_left,
            new_top=cursor_y
        )

        built_sections.append({
            "section_id": section_id,
            "label": section["label"],
            "alignment": "right" if section_id == "TITLE_OVERVIEW" else "center",
            "left_px_approx": round(section_left / EMU_PER_PX),
            "width_px_approx": round(section["width"] / EMU_PER_PX),
            "height_px_approx": round(section["height"] / EMU_PER_PX)
        })

        cursor_y += section["height"] + spacing

    replace_placeholders_on_slide(output_slide, placeholder_values or {})

    campaign_name = clean_text(placeholder_values.get("CAMPAIGN_NAME")) if placeholder_values else ""
    safe_campaign = safe_filename(campaign_name or "Campaign")
    date_stamp = datetime.now().strftime("%d%m%Y")

    filename = f"PCA One Pager_{safe_campaign}_{date_stamp}.pptx"
    output_path = GENERATED_FILES_DIR / filename

    output_prs.save(output_path)

    return {
        "output_path": str(output_path),
        "filename": filename,
        "built_sections": built_sections,
        "slide_height_px_approx": round(total_height / EMU_PER_PX)
    }



# ==================================================
# SLIDE DECK FILTERING
# ==================================================

SECTION_ID_PREFIX = "SECTION_ID:"

SECTION_ALIASES = {
    "TITLE OVERVIEW": "TITLE_OVERVIEW",
    "EXEC SUMMARY": "TITLE_OVERVIEW",
    "EXECUTIVE SUMMARY": "TITLE_OVERVIEW",
    "TITLE PERFORMANCE": "TITLE_PERFORMANCE",
    "MARKET PERFORMANCE": "MARKET_PERFORMANCE",
    "TOP TITLES MARKETS": "TOP_TITLES_MARKETS",
    "TOP TITLES & MARKETS": "TOP_TITLES_MARKETS",
    "CREATIVE OVERVIEW": "CREATIVE_OVERVIEW",
    "CREATIVE PERFORMANCE": "CREATIVE_PERFORMANCE",
    "ATTENTION SCORE": "ATTENTION_SCORE",
    "HAPPYDEMICS": "HAPPYDEMICS",
    "LUMEN RESULTS": "LUMEN_RESULTS",
    "LUMEN LEARNINGS": "LUMEN_LEARNINGS",
    "BRAND STUDY": "BRAND_STUDY",
    "LEARNING RECOMMENDATIONS": "LEARNING_RECOMMENDATIONS",
    "LEARNINGS": "LEARNING_RECOMMENDATIONS",
}


def canonical_section_id(value: Any) -> str:
    raw = clean_text(value)
    if not raw:
        return ""

    raw = raw.replace(SECTION_ID_PREFIX, "").strip()
    raw = raw.replace("&", " AND ")

    alias_key = re.sub(r"[^A-Za-z0-9]+", " ", raw).strip().upper()
    if alias_key in SECTION_ALIASES:
        return SECTION_ALIASES[alias_key]

    return re.sub(r"[^A-Za-z0-9]+", "_", raw).strip("_").upper()


def canonical_section_list(sections: List[str], include_title_overview: bool = True) -> List[str]:
    output = []
    seen = set()

    if include_title_overview:
        output.append("TITLE_OVERVIEW")
        seen.add("TITLE_OVERVIEW")

    for section in sections or []:
        section_id = canonical_section_id(section)
        if section_id and section_id not in seen:
            output.append(section_id)
            seen.add(section_id)

    return output


def iter_shapes_recursive(shapes):
    for shape in shapes:
        yield shape
        if hasattr(shape, "shapes"):
            for subshape in iter_shapes_recursive(shape.shapes):
                yield subshape


def get_shape_text_recursive(shape) -> str:
    if hasattr(shape, "text") and shape.text:
        return shape.text or ""

    if hasattr(shape, "text_frame") and shape.text_frame:
        try:
            return shape.text_frame.text or ""
        except Exception:
            pass

    return ""


def is_slide_deck_section_marker_shape(shape) -> bool:
    text = clean_text(get_shape_text_recursive(shape))
    return bool(re.match(r"^SECTION_ID\s*:", text, flags=re.IGNORECASE))


def extract_section_ids_from_slide(slide) -> List[str]:
    section_ids = []

    for shape in iter_shapes_recursive(slide.shapes):
        text = clean_text(get_shape_text_recursive(shape))
        if not text:
            continue

        matches = re.findall(
            r"SECTION_ID\s*:\s*([A-Za-z0-9_ &-]+)",
            text,
            flags=re.IGNORECASE
        )

        for match in matches:
            section_id = canonical_section_id(match)
            if section_id and section_id not in section_ids:
                section_ids.append(section_id)

    return section_ids


def remove_shape_from_slide(slide, shape) -> None:
    try:
        element = shape._element
        element.getparent().remove(element)
    except Exception:
        pass


def remove_section_id_text_boxes_from_deck(prs: Presentation) -> None:
    for slide in prs.slides:
        marker_shapes = [
            shape for shape in iter_shapes_recursive(slide.shapes)
            if is_slide_deck_section_marker_shape(shape)
        ]

        for shape in marker_shapes:
            remove_shape_from_slide(slide, shape)


def delete_slides_by_index(prs: Presentation, slide_indexes: List[int]) -> None:
    slide_id_list = prs.slides._sldIdLst

    for index in sorted(slide_indexes, reverse=True):
        if index < 0 or index >= len(prs.slides):
            continue

        slide_id = slide_id_list[index]
        prs.part.drop_rel(slide_id.rId)
        slide_id_list.remove(slide_id)


def resolve_slide_deck_sections(request: SlideDeckFromEocRequest) -> Optional[List[str]]:
    mode = clean_text(request.deck_mode).lower() or "matching"

    if mode not in ["full", "matching", "custom"]:
        raise HTTPException(status_code=400, detail="deck_mode must be full, matching, or custom.")

    if mode == "full":
        return None

    if mode == "custom":
        if not request.custom_deck_sections:
            raise HTTPException(
                status_code=400,
                detail="custom_deck_sections is required when deck_mode is custom."
            )

        return canonical_section_list(request.custom_deck_sections, include_title_overview=True)

    source_sections = request.one_pager_sections or request.selected_sections
    return canonical_section_list(source_sections, include_title_overview=True)


def filter_slide_deck_by_sections(prs: Presentation, keep_sections: Optional[List[str]]) -> Dict[str, Any]:
    if keep_sections is None:
        slide_report = []

        for index, slide in enumerate(prs.slides):
            slide_report.append({
                "slide_number": index + 1,
                "section_ids": extract_section_ids_from_slide(slide),
                "kept": True
            })

        return {
            "deck_mode": "full",
            "requested_sections": "ALL",
            "deleted_slide_count": 0,
            "kept_slide_count": len(prs.slides),
            "slide_report": slide_report,
        }

    keep_set = set(keep_sections)
    slide_indexes_to_delete = []
    slide_report = []

    for index, slide in enumerate(prs.slides):
        slide_section_ids = extract_section_ids_from_slide(slide)

        # Slides without SECTION_ID are retained as global/admin slides.
        keep_slide = True if not slide_section_ids else bool(keep_set.intersection(slide_section_ids))

        slide_report.append({
            "slide_number": index + 1,
            "section_ids": slide_section_ids,
            "kept": keep_slide
        })

        if not keep_slide:
            slide_indexes_to_delete.append(index)

    delete_slides_by_index(prs, slide_indexes_to_delete)

    return {
        "deck_mode": "filtered",
        "requested_sections": keep_sections,
        "deleted_slide_count": len(slide_indexes_to_delete),
        "kept_slide_count": len(prs.slides),
        "slide_report": slide_report,
    }


def replace_placeholders_in_presentation(prs: Presentation, placeholder_values: Dict[str, Any]) -> None:
    for slide in prs.slides:
        replace_placeholders_on_slide(slide, placeholder_values)


def build_filtered_slide_deck_ppt(
    request: SlideDeckFromEocRequest,
    placeholder_values: Dict[str, Any],
    template_path: Optional[Path] = None,
) -> Dict[str, Any]:

    tpl = template_path or SLIDE_DECK_TEMPLATE_PATH
    if not tpl.exists():
        raise HTTPException(
            status_code=500,
            detail=f"Slide deck template not found at {tpl}"
        )

    prs = Presentation(str(tpl))

    deck_sections = resolve_slide_deck_sections(request)
    filter_meta = filter_slide_deck_by_sections(prs, deck_sections)

    # Remove helper labels after filtering, so they never appear in final output.
    remove_section_id_text_boxes_from_deck(prs)

    replace_placeholders_in_presentation(prs, placeholder_values or {})

    campaign_name = clean_text(placeholder_values.get("CAMPAIGN_NAME")) if placeholder_values else ""
    safe_campaign = safe_filename(campaign_name or "Campaign")
    date_stamp = datetime.now().strftime("%d%m%Y")

    filename = f"PCA Slides_{safe_campaign}_{date_stamp}.pptx"
    output_path = GENERATED_FILES_DIR / filename

    prs.save(output_path)

    filter_meta["final_slide_count"] = len(prs.slides)

    return {
        "output_path": str(output_path),
        "filename": filename,
        "filter_meta": filter_meta,
    }


@app.get("/list-modular-sections")
def list_modular_sections():
    registry = load_section_registry()
    sections = []

    for section_id, meta in registry.items():
        sections.append({
            "section_id": section_id,
            "label": meta.get("label", section_id),
            "required": meta.get("required", False),
            "default_order": meta.get("default_order", 999)
        })

    sections.sort(key=lambda item: item.get("default_order", 999))

    return {
        "app_version": APP_VERSION,
        "sections": sections
    }


@app.post("/validate-eoc")
async def validate_eoc_endpoint(payload: FileRefsPayload):
    try:
        if not payload.openaiFileIdRefs:
            raise HTTPException(status_code=400, detail="No EOC file supplied.")

        with tempfile.TemporaryDirectory() as tmp:
            tmp_dir = Path(tmp)
            eoc_path = await download_openai_file(payload.openaiFileIdRefs[0], tmp_dir)
            sheets = read_uploaded_eoc(eoc_path)
            result = build_mapped_values(sheets, filename=eoc_path.name)

        return JSONResponse(content={
            "status": "validated",
            "app_version": APP_VERSION,
            **compact_parsed_result(result),
        })

    except HTTPException:
        raise

    except Exception as e:
        return JSONResponse(
            status_code=500,
            content={
                "error": str(e),
                "trace": traceback.format_exc(),
                "app_version": APP_VERSION,
            },
        )


@app.post("/create-modular-one-pager-download-link")
def create_modular_one_pager_download_link(request: ModularOnePagerRequest):
    try:
        result = build_grouped_stacked_modular_ppt(
            selected_sections=request.selected_sections,
            placeholder_values=request.placeholder_values or {},
            top_margin_px=request.top_margin_px,
            bottom_margin_px=request.bottom_margin_px,
            section_spacing_px=request.section_spacing_px
        )

        download_url = f"{PUBLIC_BASE_URL}/files/{Path(result['filename']).name}"

        return {
            "success": True,
            "app_version": APP_VERSION,
            "filename": result["filename"],
            "download_url": download_url,
            "built_sections": result["built_sections"],
            "slide_height_px_approx": result["slide_height_px_approx"],
            "message": "Modular one-pager generated successfully."
        }

    except HTTPException:
        raise

    except Exception as e:
        traceback.print_exc()
        raise HTTPException(
            status_code=500,
            detail={
                "message": "Failed to create modular one-pager download link",
                "error": str(e)
            }
        )


@app.post("/create-modular-one-pager-from-eoc")
async def create_modular_one_pager_from_eoc(request: ModularOnePagerFromEocRequest):
    try:
        if not request.openaiFileIdRefs:
            raise HTTPException(status_code=400, detail="No EOC file supplied.")

        with tempfile.TemporaryDirectory() as tmp:
            tmp_dir = Path(tmp)
            eoc_path = await download_openai_file(request.openaiFileIdRefs[0], tmp_dir)
            sheets = read_uploaded_eoc(eoc_path)
            parsed = build_mapped_values(sheets, filename=eoc_path.name)

        mapped_values = parsed["mapped_values"]
        mapped_values["EXEC_SUMMARY"] = resolve_exec_summary_for_ppt(request.exec_summary, mapped_values)

        result = build_grouped_stacked_modular_ppt(
            selected_sections=request.selected_sections,
            placeholder_values=mapped_values,
            top_margin_px=request.top_margin_px,
            bottom_margin_px=request.bottom_margin_px,
            section_spacing_px=request.section_spacing_px
        )

        download_url = f"{PUBLIC_BASE_URL}/files/{Path(result['filename']).name}"

        return {
            "success": True,
            "app_version": APP_VERSION,
            "filename": result["filename"],
            "download_url": download_url,
            "built_sections": result["built_sections"],
            "slide_height_px_approx": result["slide_height_px_approx"],
            "summary": compact_parsed_result(parsed),
            "message": "Modular one-pager generated from EOC successfully."
        }

    except HTTPException:
        raise

    except Exception as e:
        traceback.print_exc()
        raise HTTPException(
            status_code=500,
            detail={
                "message": "Failed to create modular one-pager from EOC",
                "error": str(e),
                "trace": traceback.format_exc(),
            }
        )


@app.post("/create-modular-one-pager-from-eoc-file-response")
async def create_modular_one_pager_from_eoc_file_response(request: ModularOnePagerFromEocRequest):
    try:
        if not request.openaiFileIdRefs:
            raise HTTPException(status_code=400, detail="No EOC file supplied.")

        with tempfile.TemporaryDirectory() as tmp:
            tmp_dir = Path(tmp)
            eoc_path = await download_openai_file(request.openaiFileIdRefs[0], tmp_dir)
            sheets = read_uploaded_eoc(eoc_path)
            parsed = build_mapped_values(sheets, filename=eoc_path.name)

        mapped_values = parsed["mapped_values"]
        mapped_values["EXEC_SUMMARY"] = resolve_exec_summary_for_ppt(request.exec_summary, mapped_values)

        result = build_grouped_stacked_modular_ppt(
            selected_sections=request.selected_sections,
            placeholder_values=mapped_values,
            top_margin_px=request.top_margin_px,
            bottom_margin_px=request.bottom_margin_px,
            section_spacing_px=request.section_spacing_px
        )

        output_path = Path(result["output_path"])

        if output_path.stat().st_size > MAX_RETURN_FILE_BYTES:
            download_url = f"{PUBLIC_BASE_URL}/files/{output_path.name}"
            return {
                "success": True,
                "app_version": APP_VERSION,
                "filename": result["filename"],
                "download_url": download_url,
                "built_sections": result["built_sections"],
                "slide_height_px_approx": result["slide_height_px_approx"],
                "summary": compact_parsed_result(parsed),
                "message": "File was too large for file-card return, so a download link was created."
            }

        encoded = base64.b64encode(output_path.read_bytes()).decode("utf-8")

        return {
            "success": True,
            "app_version": APP_VERSION,
            "summary": compact_parsed_result(parsed),
            "built_sections": result["built_sections"],
            "slide_height_px_approx": result["slide_height_px_approx"],
            "openaiFileResponse": [
                {
                    "name": result["filename"],
                    "mime_type": PPTX_MIME_TYPE,
                    "content": encoded,
                }
            ],
        }

    except HTTPException:
        raise

    except Exception as e:
        traceback.print_exc()
        raise HTTPException(
            status_code=500,
            detail={
                "message": "Failed to create modular one-pager file response",
                "error": str(e),
                "trace": traceback.format_exc(),
            }
        )



@app.post("/generate-slide-deck")
async def generate_slide_deck(request: SlideDeckFromEocRequest):
    try:
        if not request.openaiFileIdRefs:
            raise HTTPException(status_code=400, detail="No EOC file supplied.")

        with tempfile.TemporaryDirectory() as tmp:
            tmp_dir = Path(tmp)
            eoc_path = await download_openai_file(request.openaiFileIdRefs[0], tmp_dir)
            sheets = read_uploaded_eoc(eoc_path)
            parsed = build_mapped_values(sheets, filename=eoc_path.name)

        mapped_values = parsed["mapped_values"]
        mapped_values["EXEC_SUMMARY"] = resolve_exec_summary_for_ppt(request.exec_summary, mapped_values)

        result = build_filtered_slide_deck_ppt(
            request=request,
            placeholder_values=mapped_values,
        )

        download_url = f"{PUBLIC_BASE_URL}/files/{Path(result['filename']).name}"

        return {
            "success": True,
            "app_version": APP_VERSION,
            "filename": result["filename"],
            "download_url": download_url,
            "filter_meta": result["filter_meta"],
            "summary": compact_parsed_result(parsed),
            "message": "Slide deck generated successfully."
        }

    except HTTPException:
        raise

    except Exception as e:
        traceback.print_exc()
        raise HTTPException(
            status_code=500,
            detail={
                "message": "Failed to generate slide deck",
                "error": str(e),
                "trace": traceback.format_exc(),
            }
        )


@app.post("/generate-slide-deck-file-response")
async def generate_slide_deck_file_response(request: SlideDeckFromEocRequest):
    try:
        if not request.openaiFileIdRefs:
            raise HTTPException(status_code=400, detail="No EOC file supplied.")

        with tempfile.TemporaryDirectory() as tmp:
            tmp_dir = Path(tmp)
            eoc_path = await download_openai_file(request.openaiFileIdRefs[0], tmp_dir)
            sheets = read_uploaded_eoc(eoc_path)
            parsed = build_mapped_values(sheets, filename=eoc_path.name)

        mapped_values = parsed["mapped_values"]
        mapped_values["EXEC_SUMMARY"] = resolve_exec_summary_for_ppt(request.exec_summary, mapped_values)

        result = build_filtered_slide_deck_ppt(
            request=request,
            placeholder_values=mapped_values,
        )

        output_path = Path(result["output_path"])

        if output_path.stat().st_size > MAX_RETURN_FILE_BYTES:
            download_url = f"{PUBLIC_BASE_URL}/files/{output_path.name}"
            return {
                "success": True,
                "app_version": APP_VERSION,
                "filename": result["filename"],
                "download_url": download_url,
                "filter_meta": result["filter_meta"],
                "summary": compact_parsed_result(parsed),
                "message": "File was too large for file-card return, so a download link was created."
            }

        encoded = base64.b64encode(output_path.read_bytes()).decode("utf-8")

        return {
            "success": True,
            "app_version": APP_VERSION,
            "summary": compact_parsed_result(parsed),
            "filter_meta": result["filter_meta"],
            "openaiFileResponse": [
                {
                    "name": result["filename"],
                    "mime_type": PPTX_MIME_TYPE,
                    "content": encoded,
                }
            ],
        }

    except HTTPException:
        raise

    except Exception as e:
        traceback.print_exc()
        raise HTTPException(
            status_code=500,
            detail={
                "message": "Failed to generate slide deck file response",
                "error": str(e),
                "trace": traceback.format_exc(),
            }
        )


@app.get("/files/{filename}")
def get_generated_file(filename: str):
    safe_name = os.path.basename(filename)

    if not safe_name.endswith(".pptx"):
        raise HTTPException(status_code=400, detail="Only .pptx files can be downloaded.")

    file_path = GENERATED_FILES_DIR / safe_name

    if not file_path.exists():
        raise HTTPException(status_code=404, detail="Generated file not found or has expired.")

    return FileResponse(
        path=str(file_path),
        filename=safe_name,
        media_type=PPTX_MIME_TYPE
    )

# ====================================================================
# MODEL CONTEXT PROTOCOL (MCP) UNIFIED PROTOCOL INTERCEPTOR
# ====================================================================
# Refactored to route PowerPoint assembly directly to the lightweight 
# link-generation endpoint, preventing stateless buffer overloads.

@app.post("/mcp")
async def mcp_post_endpoint(request: Dict[str, Any]):
    req_id = request.get("id", 1)
    method = request.get("method", "")
    params = request.get("params", {})

    # 1. Handle Protocol Handshake Initialization Check
    if method == "initialize":
        return {
            "jsonrpc": "2.0",
            "id": req_id,
            "result": {
                "protocolVersion": "2024-11-05",
                "capabilities": {
                    "tools": {}
                },
                "serverInfo": {
                    "name": "PCA_Automation_Generator",
                    "version": APP_VERSION
                }
            }
        }

    # 2. Handle Post-Initialization Event Notification
    elif method == "notifications/initialized":
        return JSONResponse(content={})

    # 3. Handle System Capability Discovery Requests
    elif method == "tools/list":
        return {
            "jsonrpc": "2.0",
            "id": req_id,
            "result": {
                "tools": [
                    {
                        "name": "validate_eoc",
                        "description": "Validate an uploaded EOC Excel file and return mapped campaign values.",
                        "inputSchema": {
                            "type": "object",
                            "properties": {
                                "openaiFileIdRefs": {
                                    "type": "array",
                                    "items": {
                                        "type": "object",
                                        "properties": {
                                            "download_link": {"type": "string"},
                                            "name": {"type": "string"},
                                            "mime_type": {"type": "string"}
                                        }
                                    }
                                }
                            },
                            "required": ["openaiFileIdRefs"]
                        }
                    },
                    {
                        "name": "create_modular_one_pager",
                        "description": "Generate a modular PowerPoint One-Pager from verified EOC metrics.",
                        "inputSchema": {
                            "type": "object",
                            "properties": {
                                "openaiFileIdRefs": {
                                    "type": "array",
                                    "items": {
                                        "type": "object",
                                        "properties": {
                                            "download_link": {"type": "string"}
                                        }
                                    }
                                },
                                "selected_sections": {
                                    "type": "array",
                                    "items": {"type": "string"}
                                },
                                "exec_summary": {"type": "string"}
                            },
                            "required": ["openaiFileIdRefs", "selected_sections"]
                        }
                    },
                    {
                        "name": "generate_slide_deck",
                        "description": "Generate a filtered PowerPoint Slide Deck from verified EOC metrics.",
                        "inputSchema": {
                            "type": "object",
                            "properties": {
                                "openaiFileIdRefs": {
                                    "type": "array",
                                    "items": {
                                        "type": "object",
                                        "properties": {
                                            "download_link": {"type": "string"},
                                            "name": {"type": "string"},
                                            "mime_type": {"type": "string"}
                                        }
                                    }
                                },
                                "selected_sections": {
                                    "type": "array",
                                    "items": {"type": "string"}
                                },
                                "one_pager_sections": {
                                    "type": "array",
                                    "items": {"type": "string"}
                                },
                                "custom_deck_sections": {
                                    "type": "array",
                                    "items": {"type": "string"}
                                },
                                "deck_mode": {
                                    "type": "string",
                                    "enum": ["full", "matching", "custom"]
                                },
                                "exec_summary": {"type": "string"}
                            },
                            "required": ["openaiFileIdRefs", "deck_mode"]
                        }
                    }
                ]
            }
        }

    # 4. Route Runtime Tool Executions straight to Core Subroutines
    elif method == "tools/call":
        tool_name = params.get("name", "")
        arguments = params.get("arguments", {})

        if tool_name == "validate_eoc":
            try:
                payload = FileRefsPayload(**arguments)
                raw_response = await validate_eoc_endpoint(payload)
                response_text = raw_response.body.decode("utf-8")
                return {
                    "jsonrpc": "2.0",
                    "id": req_id,
                    "result": {
                        "content": [{"type": "text", "text": response_text}]
                    }
                }
            except Exception as e:
                return {
                    "jsonrpc": "2.0",
                    "id": req_id,
                    "error": {"code": -32603, "message": f"Execution error: {str(e)}"}
                }

        elif tool_name == "create_modular_one_pager":
            try:
                request_obj = ModularOnePagerFromEocRequest(**arguments)
                # Switched target to the lightweight link generator to resolve the connection reset
                raw_response = await create_modular_one_pager_from_eoc(request_obj)
                if hasattr(raw_response, "body"):
                    response_text = raw_response.body.decode("utf-8")
                else:
                    response_text = json.dumps(raw_response)
                return {
                    "jsonrpc": "2.0",
                    "id": req_id,
                    "result": {
                        "content": [{"type": "text", "text": response_text}]
                    }
                }
            except Exception as e:
                return {
                    "jsonrpc": "2.0",
                    "id": req_id,
                    "error": {"code": -32603, "message": f"Execution error: {str(e)}"}
                }

        elif tool_name == "generate_slide_deck":
            try:
                request_obj = SlideDeckFromEocRequest(**arguments)
                raw_response = await generate_slide_deck(request_obj)
                if hasattr(raw_response, "body"):
                    response_text = raw_response.body.decode("utf-8")
                else:
                    response_text = json.dumps(raw_response)
                return {
                    "jsonrpc": "2.0",
                    "id": req_id,
                    "result": {
                        "content": [{"type": "text", "text": response_text}]
                    }
                }
            except Exception as e:
                return {
                    "jsonrpc": "2.0",
                    "id": req_id,
                    "error": {"code": -32603, "message": f"Execution error: {str(e)}"}
                }

    # 5. Fallback Error Block for Unsupported Operations
    return {
        "jsonrpc": "2.0",
        "id": req_id,
        "error": {"code": -32601, "message": f"Method not found: {method}"}
    }

# ===== Web app routes =====
from web_routes import register_web_routes
register_web_routes(app)
